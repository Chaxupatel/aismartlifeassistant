import os
from pptx import Presentation
from pptx.util import Inches, Pt

def add_title_slide(prs, title, subtitle):
    slide_layout = prs.slide_layouts[0] # Title slide
    slide = prs.slides.add_slide(slide_layout)
    
    title_shape = slide.shapes.title
    subtitle_shape = slide.placeholders[1]
    
    title_shape.text = title
    subtitle_shape.text = subtitle

def add_feature_slide(prs, title, content, image_path):
    # Use Title and Content layout or a Blank layout
    slide_layout = prs.slide_layouts[5] # Title only
    slide = prs.slides.add_slide(slide_layout)
    
    # Set title
    title_shape = slide.shapes.title
    title_shape.text = title
    
    # Add text box for content
    left = Inches(0.5)
    top = Inches(2.0)
    width = Inches(4.5)
    height = Inches(4.5)
    
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    
    p = tf.add_paragraph()
    p.text = content
    p.font.size = Pt(20)
    p.line_spacing = 1.5
    
    # Add image if exists
    if os.path.exists(image_path):
        img_left = Inches(5.5)
        img_top = Inches(1.5)
        img_height = Inches(5.5)
        try:
            slide.shapes.add_picture(image_path, img_left, img_top, height=img_height)
        except Exception as e:
            print(f"Error adding image {image_path}: {e}")
    else:
        print(f"Warning: Image '{image_path}' not found. Skipping image for slide '{title}'.")
        
        # Add placeholder shape
        img_left = Inches(5.5)
        img_top = Inches(1.5)
        img_width = Inches(3.5)
        img_height = Inches(5.5)
        shape = slide.shapes.add_shape(
            1, # Rectangle
            img_left, img_top, img_width, img_height
        )
        shape.text = f"Place '{os.path.basename(image_path)}' here"

def main():
    prs = Presentation()
    
    # Define slides content and expected screenshot names
    slides_data = [
        {
            "title": "Welcome to Remindly",
            "content": "The AI-powered smart life assistant designed to keep you productive, organized, and focused.\n\n• Built with Flutter\n• Modern Bento Box UI\n• AI-driven Insights",
            "image": "home_screen.png"
        },
        {
            "title": "Interactive Bento Dashboard",
            "content": "A stunning, customizable home screen.\n\n• Dynamic greetings and weather\n• Quick access to Reminders and Events\n• Premium glassmorphism design",
            "image": "home_screen.png"
        },
        {
            "title": "Conversational AI Assistant",
            "content": "Just chat to manage your life.\n\n• Natural language parsing to schedule reminders\n• Smart productivity insights\n• Contextual suggestions presented in a 2x2 Bento grid",
            "image": "ai_assistant.png"
        },
        {
            "title": "Calendar & Events",
            "content": "A comprehensive view of your time.\n\n• Interactive monthly calendar\n• Masonry grid for upcoming events\n• Color-coded categorization for Work, Personal, and Health",
            "image": "calendar.png"
        },
        {
            "title": "Smart Reminders",
            "content": "Never miss a beat.\n\n• Beautiful 2-column masonry grid layout\n• Support for daily, weekly, and custom recurrences\n• Quick complete toggles",
            "image": "reminders.png"
        },
        {
            "title": "Profile & Settings",
            "content": "Tailored to your preferences.\n\n• Full dark mode support\n• Granular alarm and notification controls\n• Secure Firebase authentication",
            "image": "profile.png"
        }
    ]
    
    # Add title slide
    add_title_slide(prs, "Remindly: AI Smart Life Assistant", "App Overview & Features")
    
    # Add feature slides
    for slide in slides_data:
        image_path = os.path.join("screenshots", slide["image"])
        add_feature_slide(prs, slide["title"], slide["content"], image_path)
    
    # Save presentation
    output_file = "Remindly_Presentation.pptx"
    prs.save(output_file)
    print(f"Presentation saved successfully as {output_file}")
    print("If you haven't already, please add your screenshots to the 'screenshots' folder and run this script again for the best results!")

if __name__ == "__main__":
    main()
